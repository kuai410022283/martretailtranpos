from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from datetime import datetime
import os

class ExcelExporter:
    def __init__(self):
        self.wb = Workbook()
        if 'Sheet' in self.wb.sheetnames:
            del self.wb['Sheet']
    
    def create_sheet(self, sheet_name, headers, data, title=None):
        ws = self.wb.create_sheet(sheet_name)
        
        if title:
            ws.merge_cells('A1:' + chr(ord('A') + len(headers) - 1) + '1')
            title_cell = ws['A1']
            title_cell.value = title
            title_cell.font = Font(size=16, bold=True)
            title_cell.alignment = Alignment(horizontal='center', vertical='center')
            row_offset = 2
        else:
            row_offset = 1
        
        header_font = Font(bold=True)
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_offset, column=col_idx, value=header)
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border
        
        for row_idx, row_data in enumerate(data, start=row_offset + 1):
            for col_idx, cell_data in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=cell_data)
                cell.alignment = Alignment(horizontal='left', vertical='center')
                cell.border = thin_border
        
        for col_idx in range(1, len(headers) + 1):
            ws.column_dimensions[chr(ord('A') + col_idx - 1)].auto_size = True
        
        return ws
    
    def save(self, file_path):
        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        self.wb.save(file_path)
        return file_path
    
    def export_sales_report(self, file_path, sales_data, start_date, end_date):
        headers = ['销售单号', '会员', '总金额', '实付金额', '支付方式', '收银员', '时间']
        title = f'销售报表 ({start_date} 至 {end_date})'
        
        self.create_sheet('销售记录', headers, sales_data, title)
        return self.save(file_path)
    
    def export_stock_report(self, file_path, stock_data):
        headers = ['商品编码', '商品名称', '分类', '品牌', '库存数量', '成本价', '售价', '状态']
        title = '库存报表'
        
        self.create_sheet('库存信息', headers, stock_data, title)
        return self.save(file_path)
    
    def export_product_list(self, file_path, product_data):
        headers = ['商品编码', '条码', '商品名称', '分类', '品牌', '规格', '成本价', '售价', '会员价', '库存', '状态']
        title = '商品列表'
        
        self.create_sheet('商品列表', headers, product_data, title)
        return self.save(file_path)
