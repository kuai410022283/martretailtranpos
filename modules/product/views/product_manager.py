from PyQt5.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QPushButton, 
    QTableWidget, QTableWidgetItem, QHeaderView,
    QLineEdit, QComboBox, QLabel, QMessageBox, QDialog,
    QFormLayout, QDoubleSpinBox, QTextEdit, QFileDialog
)
from PyQt5.QtCore import Qt
from PyQt5.QtGui import QFont, QPixmap
from sqlalchemy.orm import Session
from shared.models.product import Product, Category
from modules.product.controllers import ProductController
from core.logger import logger
import os
import shutil
import openpyxl
from datetime import datetime

class ProductManager(QWidget):
    def __init__(self, db_session: Session, current_user):
        super().__init__()
        self.db_session = db_session
        self.current_user = current_user
        self.products = []
        self.init_ui()
        self.load_products()
    
    def init_ui(self):
        self.setWindowTitle("商品管理")
        
        # 订阅销售完成事件，自动刷新数据
        from core.event_bus import subscribe_event
        subscribe_event("sale_completed", self.on_sale_completed)
        
        layout = QVBoxLayout()
        
        toolbar = QHBoxLayout()
        
        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("搜索商品...")
        self.search_input.setFixedWidth(300)
        self.search_input.textChanged.connect(self.filter_products)
        
        self.category_combo = QComboBox()
        self.category_combo.addItem("全部分类", None)
        self.category_combo.currentIndexChanged.connect(self.filter_products)
        
        self.status_combo = QComboBox()
        self.status_combo.addItem("全部状态", None)
        self.status_combo.addItem("正常", "正常")
        self.status_combo.addItem("已停用", "已停用")
        self.status_combo.currentIndexChanged.connect(self.filter_products)
        
        add_button = QPushButton("新增商品")
        add_button.clicked.connect(self.add_product)
        
        edit_button = QPushButton("编辑商品")
        edit_button.clicked.connect(self.edit_product)
        
        delete_button = QPushButton("停用商品")
        delete_button.clicked.connect(self.delete_product)
        
        batch_enable_button = QPushButton("批量启用")
        batch_enable_button.clicked.connect(self.batch_enable_products)
        
        batch_disable_button = QPushButton("批量停用")
        batch_disable_button.clicked.connect(self.batch_disable_products)
        
        refresh_button = QPushButton("刷新")
        refresh_button.clicked.connect(self.load_products)
        
        import_button = QPushButton("批量导入")
        import_button.clicked.connect(self.import_products)
        
        category_button = QPushButton("分类管理")
        category_button.clicked.connect(self.manage_categories)
        
        toolbar.addWidget(QLabel("搜索:"))
        toolbar.addWidget(self.search_input)
        toolbar.addWidget(QLabel("分类:"))
        toolbar.addWidget(self.category_combo)
        toolbar.addWidget(QLabel("状态:"))
        toolbar.addWidget(self.status_combo)
        toolbar.addStretch()
        toolbar.addWidget(add_button)
        toolbar.addWidget(edit_button)
        toolbar.addWidget(delete_button)
        toolbar.addWidget(batch_enable_button)
        toolbar.addWidget(batch_disable_button)
        toolbar.addWidget(refresh_button)
        toolbar.addWidget(import_button)
        toolbar.addWidget(category_button)
        
        layout.addLayout(toolbar)
        
        self.table = QTableWidget()
        self.table.setColumnCount(10)
        self.table.setHorizontalHeaderLabels([
            "商品编码", "条码", "商品名称", "分类", "品牌",
            "成本价", "售价", "库存", "状态", "创建时间"
        ])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.setSelectionBehavior(QTableWidget.SelectRows)
        self.table.setAlternatingRowColors(True)
        
        layout.addWidget(self.table)
        
        self.setLayout(layout)
        
        self.load_categories()
    
    def on_sale_completed(self, **kwargs):
        """当销售完成时触发，刷新商品数据"""
        logger.info("收到销售完成事件，正在刷新商品数据...")
        self.load_products()
    
    def load_categories(self):
        try:
            with self.db_session() as db:
                controller = ProductController(db)
                categories = controller.get_categories()
                self.category_combo.clear()
                self.category_combo.addItem("全部分类", None)
                for cat in categories:
                    self.category_combo.addItem(cat.name, cat.id)
        except Exception as e:
            logger.error(f"加载分类失败: {e}")
    
    def load_products(self):
        try:
            with self.db_session() as db:
                controller = ProductController(db)
                keyword = self.search_input.text().strip()
                category_id = self.category_combo.currentData()
                status = self.status_combo.currentData()
                
                self.products = controller.list_products(
                    category_id=category_id,
                    status=status,
                    keyword=keyword if keyword else None
                )
                
                self.table.setRowCount(len(self.products))
                
                for row, product in enumerate(self.products):
                    self.table.setItem(row, 0, QTableWidgetItem(product.code))
                    self.table.setItem(row, 1, QTableWidgetItem(product.barcode or ""))
                    self.table.setItem(row, 2, QTableWidgetItem(product.name))
                    
                    category_name = product.category.name if product.category else ""
                    self.table.setItem(row, 3, QTableWidgetItem(category_name))
                    
                    self.table.setItem(row, 4, QTableWidgetItem(product.brand or ""))
                    self.table.setItem(row, 5, QTableWidgetItem(f"{product.price_cost:.2f}"))
                    self.table.setItem(row, 6, QTableWidgetItem(f"{product.price_sale:.2f}"))
                    self.table.setItem(row, 7, QTableWidgetItem(f"{product.stock_quantity:.2f}"))
                    self.table.setItem(row, 8, QTableWidgetItem(product.status))
                    self.table.setItem(row, 9, QTableWidgetItem(product.create_time.strftime("%Y-%m-%d %H:%M")))
                    
                    for col in range(10):
                        self.table.item(row, col).setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)
        
        except Exception as e:
            logger.error(f"加载商品列表失败: {e}")
            QMessageBox.critical(self, "错误", f"加载商品列表失败: {str(e)}")
    
    def filter_products(self):
        self.load_products()
    
    def add_product(self):
        dialog = ProductEditDialog(self.db_session, self, None)
        if dialog.exec_() == QDialog.Accepted:
            self.load_products()
            QMessageBox.information(self, "成功", "商品添加成功!")
    
    def edit_product(self):
        selected_rows = self.table.selectionModel().selectedRows()
        if not selected_rows:
            QMessageBox.warning(self, "提示", "请先选择要编辑的商品")
            return
        
        row = selected_rows[0].row()
        product = self.products[row]
        
        dialog = ProductEditDialog(self.db_session, self, product)
        if dialog.exec_() == QDialog.Accepted:
            self.load_products()
            QMessageBox.information(self, "成功", "商品更新成功!")
    
    def delete_product(self):
        selected_rows = self.table.selectionModel().selectedRows()
        if not selected_rows:
            QMessageBox.warning(self, "提示", "请先选择要停用的商品")
            return
        
        row = selected_rows[0].row()
        product = self.products[row]
        
        reply = QMessageBox.question(
            self, "确认",
            f"确定要停用商品 {product.name} 吗？",
            QMessageBox.Yes | QMessageBox.No
        )
        
        if reply == QMessageBox.Yes:
            try:
                with self.db_session() as db:
                    controller = ProductController(db)
                    if controller.delete_product(product.id):
                        self.load_products()
                        QMessageBox.information(self, "成功", "商品已停用")
            except Exception as e:
                logger.error(f"停用商品失败: {e}")
                QMessageBox.critical(self, "错误", f"停用商品失败: {str(e)}")
    
    def batch_enable_products(self):
        selected_rows = self.table.selectionModel().selectedRows()
        if not selected_rows:
            QMessageBox.warning(self, "提示", "请先选择要启用的商品")
            return
        
        product_ids = [self.products[row.row()].id for row in selected_rows]
        
        reply = QMessageBox.question(
            self, "确认", f"确定要启用选中的 {len(product_ids)} 个商品吗？",
            QMessageBox.Yes | QMessageBox.No
        )
        
        if reply == QMessageBox.Yes:
            try:
                with self.db_session() as db:
                    controller = ProductController(db)
                    for product_id in product_ids:
                        controller.update_product(product_id, status="正常")
                self.load_products()
                QMessageBox.information(self, "成功", "商品已批量启用")
            except Exception as e:
                logger.error(f"批量启用商品失败: {e}")
                QMessageBox.critical(self, "错误", f"批量启用商品失败: {str(e)}")
    
    def batch_disable_products(self):
        selected_rows = self.table.selectionModel().selectedRows()
        if not selected_rows:
            QMessageBox.warning(self, "提示", "请先选择要停用的商品")
            return
        
        product_ids = [self.products[row.row()].id for row in selected_rows]
        
        reply = QMessageBox.question(
            self, "确认", f"确定要停用选中的 {len(product_ids)} 个商品吗？",
            QMessageBox.Yes | QMessageBox.No
        )
        
        if reply == QMessageBox.Yes:
            try:
                with self.db_session() as db:
                    controller = ProductController(db)
                    for product_id in product_ids:
                        controller.update_product(product_id, status="已停用")
                self.load_products()
                QMessageBox.information(self, "成功", "商品已批量停用")
            except Exception as e:
                logger.error(f"批量停用商品失败: {e}")
                QMessageBox.critical(self, "错误", f"批量停用商品失败: {str(e)}")
    
    def import_products(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "选择导入文件", "", "Excel Files (*.xlsx *.xls)"
        )
        if not file_path:
            return
            
        try:
            wb = openpyxl.load_workbook(file_path)
            ws = wb.active
            
            headers = [cell.value for cell in ws[1]]
            required_headers = ['商品编码', '商品名称', '售价']
            
            # Check headers
            header_map = {}
            for i, h in enumerate(headers):
                if h:
                    header_map[h] = i
            
            missing = [h for h in required_headers if h not in header_map]
            if missing:
                 QMessageBox.warning(self, "错误", f"缺少必要列: {', '.join(missing)}")
                 return
            
            success_count = 0
            fail_count = 0
            errors = []
            
            with self.db_session() as db:
                controller = ProductController(db)
                
                # Get category map for quick lookup
                categories = controller.get_categories()
                category_map = {c.name: c.id for c in categories}
                
                for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    try:
                        code = row[header_map['商品编码']]
                        if not code: continue
                        
                        # Category lookup
                        cat_name = row[header_map.get('分类')] if '分类' in header_map else None
                        cat_id = category_map.get(cat_name) if cat_name else None
                        
                        product_data = {
                            'code': str(code),
                            'name': row[header_map['商品名称']],
                            'price_sale': float(row[header_map['售价']] or 0),
                            'barcode': str(row[header_map.get('条码')]) if '条码' in header_map and row[header_map['条码']] else None,
                            'category_id': cat_id,
                            'brand': row[header_map.get('品牌')] if '品牌' in header_map else None,
                            'specification': str(row[header_map.get('规格')]) if '规格' in header_map and row[header_map['规格']] else None,
                            'price_cost': float(row[header_map.get('成本价')] or 0) if '成本价' in header_map else 0,
                            'price_member': float(row[header_map.get('会员价')] or 0) if '会员价' in header_map else 0,
                            'stock_quantity': float(row[header_map.get('库存')] or 0) if '库存' in header_map else 0,
                            'min_stock': float(row[header_map.get('最低库存')] or 0) if '最低库存' in header_map else 0,
                            'status': '正常'
                        }
                        
                        # Check if exists
                        existing = controller.get_product_by_code(product_data['code'])
                        if existing:
                            controller.update_product(existing.id, **product_data)
                        else:
                            controller.create_product(**product_data)
                        
                        success_count += 1
                        
                    except Exception as e:
                        fail_count += 1
                        errors.append(f"行 {row_idx}: {str(e)}")
            
            msg = f"导入完成\n成功: {success_count}\n失败: {fail_count}"
            if errors:
                msg += "\n\n错误详情(前10条):\n" + "\n".join(errors[:10])
                
            QMessageBox.information(self, "导入结果", msg)
            self.load_products()
            
        except Exception as e:
            logger.error(f"导入失败: {e}")
            QMessageBox.critical(self, "错误", f"导入失败: {str(e)}")
    
    def manage_categories(self):
        dialog = CategoryManagerDialog(self.db_session, self)
        if dialog.exec_() == QDialog.Accepted:
            self.load_categories()
            self.load_products()


class CategoryManagerDialog(QDialog):
    def __init__(self, db_session, parent):
        super().__init__(parent)
        self.db_session = db_session
        self.categories = []
        self.init_ui()
        self.load_categories()
    
    def init_ui(self):
        self.setWindowTitle("商品分类管理")
        self.setMinimumWidth(500)
        
        layout = QVBoxLayout()
        
        toolbar = QHBoxLayout()
        add_button = QPushButton("新增分类")
        add_button.clicked.connect(self.add_category)
        edit_button = QPushButton("编辑分类")
        edit_button.clicked.connect(self.edit_category)
        delete_button = QPushButton("删除分类")
        delete_button.clicked.connect(self.delete_category)
        toolbar.addStretch()
        toolbar.addWidget(add_button)
        toolbar.addWidget(edit_button)
        toolbar.addWidget(delete_button)
        
        self.category_table = QTableWidget()
        self.category_table.setColumnCount(2)
        self.category_table.setHorizontalHeaderLabels(["分类名称", "创建时间"])
        self.category_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.category_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.category_table.setAlternatingRowColors(True)
        
        layout.addLayout(toolbar)
        layout.addWidget(self.category_table)
        
        buttons = QHBoxLayout()
        close_button = QPushButton("关闭")
        close_button.clicked.connect(self.reject)
        buttons.addStretch()
        buttons.addWidget(close_button)
        
        layout.addLayout(buttons)
        self.setLayout(layout)
    
    def load_categories(self):
        try:
            with self.db_session() as db:
                controller = ProductController(db)
                self.categories = controller.get_categories()
                self.category_table.setRowCount(len(self.categories))
                
                for row, category in enumerate(self.categories):
                    self.category_table.setItem(row, 0, QTableWidgetItem(category.name))
                    self.category_table.setItem(row, 1, QTableWidgetItem(category.create_time.strftime("%Y-%m-%d %H:%M")))
                    
                    for col in range(2):
                        self.category_table.item(row, col).setFlags(Qt.ItemIsSelectable | Qt.ItemIsEnabled)
        except Exception as e:
            logger.error(f"加载分类失败: {e}")
    
    def add_category(self):
        dialog = CategoryEditDialog(self.db_session, self, None)
        if dialog.exec_() == QDialog.Accepted:
            self.load_categories()
    
    def edit_category(self):
        selected_rows = self.category_table.selectionModel().selectedRows()
        if not selected_rows:
            QMessageBox.warning(self, "提示", "请先选择要编辑的分类")
            return
        
        row = selected_rows[0].row()
        category = self.categories[row]
        
        dialog = CategoryEditDialog(self.db_session, self, category)
        if dialog.exec_() == QDialog.Accepted:
            self.load_categories()
    
    def delete_category(self):
        selected_rows = self.category_table.selectionModel().selectedRows()
        if not selected_rows:
            QMessageBox.warning(self, "提示", "请先选择要删除的分类")
            return
        
        row = selected_rows[0].row()
        category = self.categories[row]
        
        reply = QMessageBox.question(
            self, "确认", f"确定要删除分类 {category.name} 吗？",
            QMessageBox.Yes | QMessageBox.No
        )
        
        if reply == QMessageBox.Yes:
            try:
                with self.db_session() as db:
                    controller = ProductController(db)
                    if controller.delete_category(category.id):
                        self.load_categories()
                        QMessageBox.information(self, "成功", "分类已删除")
            except Exception as e:
                logger.error(f"删除分类失败: {e}")
                QMessageBox.critical(self, "错误", f"删除分类失败: {str(e)}")


class CategoryEditDialog(QDialog):
    def __init__(self, db_session, parent, category=None):
        super().__init__(parent)
        self.db_session = db_session
        self.category = category
        self.init_ui()
        
        if category:
            self.setWindowTitle("编辑分类")
            self.load_category_data()
        else:
            self.setWindowTitle("新增分类")
    
    def init_ui(self):
        self.setMinimumWidth(400)
        
        layout = QFormLayout()
        
        self.name_input = QLineEdit()
        self.desc_input = QTextEdit()
        self.desc_input.setMaximumHeight(80)
        
        layout.addRow("分类名称 *:", self.name_input)
        layout.addRow("备注:", self.desc_input)
        
        buttons = QHBoxLayout()
        ok_button = QPushButton("确定")
        ok_button.clicked.connect(self.save_category)
        cancel_button = QPushButton("取消")
        cancel_button.clicked.connect(self.reject)
        buttons.addStretch()
        buttons.addWidget(ok_button)
        buttons.addWidget(cancel_button)
        
        main_layout = QVBoxLayout()
        main_layout.addLayout(layout)
        main_layout.addLayout(buttons)
        self.setLayout(main_layout)
    
    def load_category_data(self):
        self.name_input.setText(self.category.name)
        self.desc_input.setText(self.category.description or "")
    
    def save_category(self):
        name = self.name_input.text().strip()
        
        if not name:
            QMessageBox.warning(self, "提示", "请输入分类名称")
            return
        
        try:
            with self.db_session() as db:
                controller = ProductController(db)
                data = {
                    'name': name,
                    'description': self.desc_input.toPlainText().strip() or None
                }
                
                if self.category:
                    controller.update_category(self.category.id, **data)
                else:
                    controller.create_category(**data)
            
            self.accept()
        except Exception as e:
            logger.error(f"保存分类失败: {e}")
            QMessageBox.critical(self, "错误", f"保存分类失败: {str(e)}")


class ProductEditDialog(QDialog):
    def __init__(self, db_session, parent, product=None):
        super().__init__(parent)
        self.db_session = db_session
        self.product = product
        self.image_path = None
        self.init_ui()
        
        if product:
            self.setWindowTitle("编辑商品")
            self.load_product_data()
        else:
            self.setWindowTitle("新增商品")
    
    def init_ui(self):
        self.setMinimumWidth(600)
        
        main_layout = QVBoxLayout()
        
        # Image area
        image_layout = QHBoxLayout()
        self.image_label = QLabel()
        self.image_label.setFixedSize(150, 150)
        self.image_label.setStyleSheet("border: 1px solid #ccc; background-color: #f0f0f0;")
        self.image_label.setAlignment(Qt.AlignCenter)
        self.image_label.setText("无图片")
        self.image_label.setScaledContents(True)
        
        select_image_btn = QPushButton("选择图片")
        select_image_btn.clicked.connect(self.select_image)
        
        image_control_layout = QVBoxLayout()
        image_control_layout.addWidget(select_image_btn)
        image_control_layout.addStretch()
        
        image_layout.addWidget(self.image_label)
        image_layout.addLayout(image_control_layout)
        image_layout.addStretch()
        
        main_layout.addLayout(image_layout)
        
        layout = QFormLayout()
        
        self.code_input = QLineEdit()
        self.barcode_input = QLineEdit()
        self.name_input = QLineEdit()
        
        self.category_combo = QComboBox()
        
        self.brand_input = QLineEdit()
        self.spec_input = QLineEdit()
        
        self.price_cost_input = QDoubleSpinBox()
        self.price_cost_input.setRange(0, 999999)
        self.price_cost_input.setDecimals(2)
        
        self.price_sale_input = QDoubleSpinBox()
        self.price_sale_input.setRange(0, 999999)
        self.price_sale_input.setDecimals(2)
        
        self.price_member_input = QDoubleSpinBox()
        self.price_member_input.setRange(0, 999999)
        self.price_member_input.setDecimals(2)
        
        self.stock_input = QDoubleSpinBox()
        self.stock_input.setRange(0, 999999)
        self.stock_input.setDecimals(2)
        
        self.min_stock_input = QDoubleSpinBox()
        self.min_stock_input.setRange(0, 999999)
        self.min_stock_input.setDecimals(2)
        
        self.status_combo = QComboBox()
        self.status_combo.addItems(["正常", "已停用"])
        
        self.desc_input = QTextEdit()
        self.desc_input.setMaximumHeight(100)
        
        layout.addRow("商品编码 *:", self.code_input)
        layout.addRow("条码:", self.barcode_input)
        layout.addRow("商品名称 *:", self.name_input)
        layout.addRow("分类:", self.category_combo)
        layout.addRow("品牌:", self.brand_input)
        layout.addRow("规格:", self.spec_input)
        layout.addRow("成本价:", self.price_cost_input)
        layout.addRow("售价 *:", self.price_sale_input)
        layout.addRow("会员价:", self.price_member_input)
        layout.addRow("库存数量:", self.stock_input)
        layout.addRow("最低库存:", self.min_stock_input)
        layout.addRow("状态:", self.status_combo)
        layout.addRow("备注:", self.desc_input)
        
        main_layout.addLayout(layout)
        
        buttons = QHBoxLayout()
        
        ok_button = QPushButton("确定")
        ok_button.clicked.connect(self.save_product)
        
        cancel_button = QPushButton("取消")
        cancel_button.clicked.connect(self.reject)
        
        buttons.addStretch()
        buttons.addWidget(ok_button)
        buttons.addWidget(cancel_button)
        
        main_layout.addLayout(buttons)
        
        self.setLayout(main_layout)
        
        self.load_categories()
    
    def select_image(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "选择商品图片", "", "Images (*.png *.jpg *.jpeg *.bmp)"
        )
        if file_path:
            self.image_path = file_path
            pixmap = QPixmap(file_path)
            if not pixmap.isNull():
                self.image_label.setPixmap(pixmap)
            else:
                self.image_label.setText("无效图片")

    def load_categories(self):
        try:
            with self.db_session() as db:
                controller = ProductController(db)
                categories = controller.get_categories()
                self.category_combo.clear()
                self.category_combo.addItem("请选择分类", None)
                for cat in categories:
                    self.category_combo.addItem(cat.name, cat.id)
        except Exception as e:
            logger.error(f"加载分类失败: {e}")
    
    def load_product_data(self):
        self.code_input.setText(self.product.code)
        self.code_input.setReadOnly(True)
        self.barcode_input.setText(self.product.barcode or "")
        self.name_input.setText(self.product.name)
        
        index = self.category_combo.findData(self.product.category_id)
        if index >= 0:
            self.category_combo.setCurrentIndex(index)
        
        self.brand_input.setText(self.product.brand or "")
        self.spec_input.setText(self.product.specification or "")
        self.price_cost_input.setValue(self.product.price_cost)
        self.price_sale_input.setValue(self.product.price_sale)
        self.price_member_input.setValue(self.product.price_member)
        self.stock_input.setValue(self.product.stock_quantity)
        self.min_stock_input.setValue(self.product.min_stock)
        
        index = self.status_combo.findText(self.product.status)
        if index >= 0:
            self.status_combo.setCurrentIndex(index)
        
        self.desc_input.setText(self.product.description or "")
        
        # Load image
        if self.product.image_path and os.path.exists(self.product.image_path):
            self.image_path = self.product.image_path
            self.image_label.setPixmap(QPixmap(self.image_path))
    
    def save_product(self):
        code = self.code_input.text().strip()
        name = self.name_input.text().strip()
        price_sale = self.price_sale_input.value()
        
        if not code:
            QMessageBox.warning(self, "提示", "请输入商品编码")
            return
        
        if not name:
            QMessageBox.warning(self, "提示", "请输入商品名称")
            return
        
        if price_sale <= 0:
            QMessageBox.warning(self, "提示", "请输入有效的售价")
            return
        
        try:
            # Handle image
            final_image_path = self.product.image_path if self.product else None
            
            if self.image_path and self.image_path != final_image_path:
                # Copy to resources
                resources_dir = os.path.join(os.getcwd(), "resources", "products")
                if not os.path.exists(resources_dir):
                    os.makedirs(resources_dir)
                
                ext = os.path.splitext(self.image_path)[1]
                new_filename = f"{code}_{datetime.now().strftime('%Y%m%d%H%M%S')}{ext}"
                target_path = os.path.join(resources_dir, new_filename)
                
                try:
                    shutil.copy2(self.image_path, target_path)
                    final_image_path = target_path
                except Exception as e:
                    logger.error(f"复制图片失败: {e}")
                    # Fallback to original path if copy fails
                    final_image_path = self.image_path
            
            with self.db_session() as db:
                controller = ProductController(db)
                data = {
                    'code': code,
                    'barcode': self.barcode_input.text().strip() or None,
                    'name': name,
                    'category_id': self.category_combo.currentData(),
                    'brand': self.brand_input.text().strip() or None,
                    'specification': self.spec_input.text().strip() or None,
                    'price_cost': self.price_cost_input.value(),
                    'price_sale': price_sale,
                    'price_member': self.price_member_input.value(),
                    'stock_quantity': self.stock_input.value(),
                    'min_stock': self.min_stock_input.value(),
                    'status': self.status_combo.currentText(),
                    'description': self.desc_input.toPlainText().strip() or None,
                    'image_path': final_image_path
                }
                
                if self.product:
                    controller.update_product(self.product.id, **data)
                else:
                    controller.create_product(**data)
            
            self.accept()
        
        except Exception as e:
            logger.error(f"保存商品失败: {e}")
            QMessageBox.critical(self, "错误", f"保存商品失败: {str(e)}")
